import re, time
from pathlib import Path
from urllib.parse import urlparse
import pandas as pd
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / 'Retailers_with_Websites.xlsx'
OUTPUT = ROOT / 'MAP_Scan_Results.xlsx'
EVIDENCE = ROOT / 'evidence'
MAX_LINKS_PER_SITE = 5
PRODUCTS = [
 ('Distillate Cart 1g',6.20),('Distillate Cart 2g',10.08),('Plus+ 1g',7.75),('Plus+ 2g',12.40),
 ('Plus+ Pro 2g',12.40),('Flip Liquid Diamond 2g',15.50),('Live Sauce 2g',15.50),
 ('Max Pro Liquid Diamond 3g',20.93),('Gummies 200mg',2.33),('Moon Rocks 5g',19.38),
 ('Infused Single 1.25g',7.75),('Live Sauce Single 1.25g',7.75),('Max Infused Single 3g',12.40),
 ('Infused Minis 3.5g',19.38),('Live Sauce Minis 3.5g',19.38),('Infused Blunts 2-Pack 4g',15.50)]
PROMOS = ['2 for','3 for','4 for','6 for','14 for','bundle','mix and match','special','deal','sale','discount','promo']

def normalize_url(v):
    if pd.isna(v) or str(v).strip().lower() in ('','nan','none'): return ''
    v=str(v).strip()
    return v if v.startswith(('http://','https://')) else 'https://'+v

def same_domain(a,b):
    try:
        return urlparse(a).netloc.lower().removeprefix('www.') == urlparse(b).netloc.lower().removeprefix('www.')
    except Exception: return False

def main():
    print('MITTEN MAP SCAN — candidate evidence only')
    if not INPUT.exists():
        print('ERROR: Retailers_with_Websites.xlsx ko is folder mein rakho.'); input('Enter dabao...'); return
    df=pd.read_excel(INPUT)
    cols={str(c).strip().lower():c for c in df.columns}
    website_col=next((cols[k] for k in cols if k in ('website','websites','url','website url')),None)
    parent_col=next((cols[k] for k in cols if 'parent retailer' in k),None)
    customer_col=next((cols[k] for k in cols if k in ('customer','customer / location','location')),None)
    if not website_col:
        print('Website column nahi mili. Headers:',list(df.columns)); input('Enter dabao...'); return
    df['_url']=df[website_col].map(normalize_url)
    sites=df[df['_url']!=''].drop_duplicates('_url')
    EVIDENCE.mkdir(exist_ok=True)
    page_rows=[]; product_rows=[]
    print(f'Retailer rows: {len(df)} | Unique websites: {len(sites)}')
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=True)
        page=browser.new_page(viewport={'width':1440,'height':1000},ignore_https_errors=True)
        for n,(_,r) in enumerate(sites.iterrows(),1):
            url=r['_url']; parent=str(r[parent_col]) if parent_col else ''; customer=str(r[customer_col]) if customer_col else ''
            slug=re.sub(r'[^A-Za-z0-9_-]+','_',urlparse(url).netloc)[:50] or f'site{n}'
            print(f'[{n}/{len(sites)}] {url}')
            try:
                page.goto(url,wait_until='domcontentloaded',timeout=20000); page.wait_for_timeout(1200)
                links=page.locator('a[href]').evaluate_all("els=>els.map(a=>({url:a.href,text:(a.innerText||a.getAttribute('aria-label')||'').trim()}))")
                urls=[url]; seen={url}
                terms=('mitten','product','shop','menu','deal','special','sale','search','disposable','flower','gummy','pre-roll')
                for link in links:
                    href=link.get('url',''); label=(link.get('text','')+' '+href).lower()
                    if href and href not in seen and same_domain(url,href) and any(t in label for t in terms):
                        urls.append(href); seen.add(href)
                    if len(urls)>=MAX_LINKS_PER_SITE+1: break
                for i,u in enumerate(urls):
                    try:
                        page.goto(u,wait_until='domcontentloaded',timeout=18000); page.wait_for_timeout(800)
                        text=page.locator('body').inner_text(timeout=7000); low=text.lower()
                        title=page.title()
                        shot=''
                        try:
                            shot=f'evidence/{slug}_page{i}.png'
                            page.screenshot(path=str(ROOT/shot),full_page=True,timeout=10000)
                        except Exception: shot=''
                        hits=[name for name,_ in PRODUCTS if name.lower() in low]
                        prices=list(dict.fromkeys(x.replace(' ','') for x in re.findall(r'\$\s?\d{1,4}(?:\.\d{2})?',text)))[:80]
                        promo=[term for term in PROMOS if term in low]
                        if 'mitten' in low or hits or promo:
                            page_rows.append({'Parent Retailer':parent,'Customer / Location':customer,'Website':url,'Page URL':u,'Page Title':title,'Product Name Matches':'; '.join(hits),'Candidate Dollar Prices':'; '.join(prices),'Promotion Terms Found':'; '.join(promo),'Screenshot':shot,'Review Status':'Needs human review','Text Snippet':re.sub(r'\s+',' ',text)[:900]})
                        for name,map_price in PRODUCTS:
                            if name.lower() in low:
                                product_rows.append({'Parent Retailer':parent,'Customer / Location':customer,'Website':url,'Product':name,'MAP ($)':map_price,'Page URL':u,'Page Title':title,'Candidate Dollar Prices':'; '.join(prices),'Screenshot':shot,'Audit Result':'Needs review — price not matched automatically','Notes':'Verify exact product, package size, advertised price and promotion manually.'})
                    except Exception as e:
                        page_rows.append({'Parent Retailer':parent,'Customer / Location':customer,'Website':url,'Page URL':u,'Review Status':f'Page error: {str(e)[:120]}'})
                    time.sleep(0.6)
            except Exception as e:
                page_rows.append({'Parent Retailer':parent,'Customer / Location':customer,'Website':url,'Page URL':url,'Review Status':f'Could not open: {str(e)[:150]}'})
            time.sleep(0.8)
        browser.close()
    with pd.ExcelWriter(OUTPUT,engine='openpyxl') as writer:
        pd.DataFrame(page_rows).to_excel(writer,sheet_name='Pages to Review',index=False)
        pd.DataFrame(product_rows).to_excel(writer,sheet_name='Product Matches',index=False)
        pd.DataFrame([{'Retailer rows':len(df),'Unique websites':len(sites),'Relevant page rows':len(page_rows),'Product-name matches':len(product_rows),'Warning':'Automated scan is not a final compliance decision.'}]).to_excel(writer,sheet_name='Summary',index=False)
        from openpyxl.styles import Font,PatternFill,Alignment
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:
                c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='17365D'); c.alignment=Alignment(wrap_text=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(max(max(len(str(x.value or '')) for x in list(col)[:80])+2,14),42)
    print('\nDone. Output:',OUTPUT); print('Screenshots:',EVIDENCE); print('Review all candidate prices; no site was auto-marked compliant/violation.')
    input('Enter dabao band karne ke liye...')

if __name__=='__main__': main()
